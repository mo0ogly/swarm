"""Build the French LinkedIn DOCX from the unabridged supplied article.

Run with the bundled document Python runtime and python-docx. The builder reads
article-fr.md and accompagnement-original-fr.md, preserves their prose and
headings, performs no network access and atomically replaces its sole DOCX output.
Rendering and visual approval are separate steps; this program proves neither.
"""

from __future__ import annotations

import argparse
import io
import os
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm, Pt, RGBColor


def add_inline(paragraph, text: str) -> None:
    """Append original prose with Markdown links, emphasis or inline code.

Only HTTPS references and non-nested emphasis/code spans are supported. Other text is
literal. Invalid link schemes raise ValueError before the output is written;
links create external DOCX relationships but are never fetched by this builder.
"""
    tokens = re.compile(r"(\[[^\]]+\]\([^\s)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    for token in tokens.split(text):
        if not token:
            continue
        if token.startswith("[") and "](" in token:
            match = re.fullmatch(r"\[([^\]]+)\]\(([^\s)]+)\)", token)
            if match is None:
                raise ValueError("Malformed editorial hyperlink")
            label, url = match.groups()
            url = url.replace("\\&", "&")
            if not url.startswith("https://"):
                raise ValueError("Only HTTPS references are accepted")
            relationship = paragraph.part.relate_to(
                url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relationship)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "24588A")
            underline = OxmlElement("w:u")
            underline.set(qn("w:val"), "single")
            properties.extend([color, underline])
            run.append(properties)
            value = OxmlElement("w:t")
            value.text = label
            run.append(value)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`") and token.endswith("`"):
            paragraph.add_run(token[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(token)


def add_article(document, source: str) -> None:
    """Append the supported article Markdown without omitting source paragraphs.

    Input accepts headings, prose, bullets, numbered lines and Markdown tables.
Original titles and punctuation are retained at the user's explicit request.
Fenced code and HTML raise ValueError instead of being dropped. Mutates the in-memory
document only; unsupported structure leaves an existing output untouched.
"""
    for block in source.strip().split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block == "---":
            document.add_paragraph()
            continue
        if block.startswith(("```", "<")):
            raise ValueError("Unsupported article structure")
        heading = re.match(r"^(#{1,6}) (.+)$", block)
        if heading:
            level, title = len(heading[1]), heading[2]
            style = "Title" if level == 1 else f"Heading {level-1}"
            add_inline(document.add_paragraph(style=style), title)
            continue
        lines = block.splitlines()
        if block.startswith("|"):
            rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
            rows = [row for row in rows if not all(re.fullmatch(r":?-+:?", cell) for cell in row)]
            if not rows or any(len(row) != len(rows[0]) for row in rows):
                raise ValueError("Inconsistent source table")
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            borders = OxmlElement("w:tblBorders")
            for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                border = OxmlElement("w:" + edge)
                for name, value in {"val":"single", "sz":"4", "color":"D9D9D9"}.items():
                    border.set(qn("w:" + name), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for index, cells in enumerate(rows):
                row = table.add_row()
                for cell, text in zip(row.cells, cells):
                    add_inline(cell.paragraphs[0], text)
                    if index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "EAEAEA")
                        cell._tc.get_or_add_tcPr().append(shading)
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        if all(line.startswith("- ") for line in lines):
            for line in lines:
                add_inline(document.add_paragraph(style="List Bullet"), line[2:])
            continue
        if all(re.match(r"^\d+\. ", line) for line in lines):
            for line in lines:
                add_inline(document.add_paragraph(style="Reference"), line)
        else:
            add_inline(document.add_paragraph(), " ".join(lines))


def save_atomically(document, output: Path) -> None:
    """Save a checked DOCX through a sibling temporary file and atomic replace.

The destination directory must exist. Serialization, ZIP validation, fsync and
replacement errors propagate with their original causes. Before replacement an
existing output remains intact; a failed temporary file is retained for diagnosis.
A directory fsync failure after replacement means publication may have occurred.
"""
    stream = io.BytesIO()
    document.save(stream)
    payload = stream.getvalue()
    with ZipFile(io.BytesIO(payload)) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError("DOCX contains an invalid ZIP member: " + bad_member)
        if "word/document.xml" not in archive.namelist():
            raise ValueError("DOCX is missing its document body")
    descriptor, temporary = tempfile.mkstemp(
        prefix="." + output.stem + "-", suffix=".docx", dir=output.parent
    )
    with os.fdopen(descriptor, "wb") as handle:
        handle.write(payload)
        handle.flush()
        os.fsync(handle.fileno())
    if Path(temporary).read_bytes() != payload:
        raise OSError("Temporary DOCX verification failed")
    os.replace(temporary, output)
    directory = os.open(output.parent, os.O_RDONLY | os.O_DIRECTORY)
    try:
        os.fsync(directory)
    finally:
        os.close(directory)


def build(source_dir: Path, output: Path) -> None:
    """Create the original article followed by its supplied accompanying notes.

Sources must exist as UTF-8. Uses A4 pages, black native Word headings and
    editable paragraphs and the original notes table; no images, active content
    or credentials are added. The original text, including punctuation, is retained.
Reading, parsing and atomic save errors propagate. No render approval is implied.
"""
    article = (source_dir / "article-fr.md").read_text(encoding="utf-8")
    accompaniment = (source_dir / "accompagnement-original-fr.md").read_text(encoding="utf-8")
    document = Document()
    # The bundled default template carries a blue Title paragraph border.
    # Remove paragraph borders at their source so inherited formatting cannot
    # reintroduce a decorative rule in Word or the headless renderer.
    for border in list(document.styles.element.xpath(".//w:pBdr")):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.1)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.widow_control = True
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.keep_together = True
    title = document.styles["Title"]
    title.font.size = Pt(24)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(16)
    title.paragraph_format.keep_with_next = True
    heading = document.styles["Heading 1"]
    heading.font.size = Pt(14)
    heading.font.bold = True
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True
    document.styles["Heading 2"].font.size = Pt(12)
    document.styles["Heading 2"].font.bold = True
    document.styles["Heading 2"].paragraph_format.keep_with_next = True
    reference = document.styles.add_style("Reference", 1)
    reference.base_style = normal
    reference.font.size = Pt(10)
    reference.paragraph_format.space_after = Pt(10)
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "fr-FR")
    normal.element.get_or_add_rPr().append(language)
    document.core_properties.title = "IA agentique : qui garde les règles, les clés et les conséquences ?"
    document.core_properties.subject = "Positionnement et références de sécurité"
    document.core_properties.author = "Fabrice Pizzi"
    document.core_properties.keywords = "Swarm, sécurité agentique, LinkedIn"
    add_article(document, article)
    document.add_page_break()
    document.add_paragraph("Annexe — Éléments d’accompagnement du texte fourni", "Title")
    add_article(document, accompaniment)
    save_atomically(document, output)
    print(f"Created {output.name}; original article: {len(article.split())} words; accompanying notes: {len(accompaniment.split())} words")


def main() -> None:
    """Parse local source/output paths and build one DOCX, propagating failures."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=Path(__file__).parent)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.source_dir / "Swarm_Securite_LinkedIn.docx"
    build(args.source_dir, output)


if __name__ == "__main__":
    main()
