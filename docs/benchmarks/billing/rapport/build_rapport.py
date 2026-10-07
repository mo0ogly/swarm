#!/usr/bin/env python3
"""Rapport Word des résultats du banc de facturation : campagne scriptée et lot à agents réels.

Usage : build_rapport.py --scripted CAMPAGNE.jsonl [--real LOT.jsonl] --out RAPPORT.docx

Toutes les valeurs viennent des fichiers JSONL, par les modules d'analyse du banc (tables,
hypotheses, tables_reel) : le rapport n'ajoute aucun chiffre. Une campagne sans ligne de fin est
signalée comme provisoire. Nécessite python-docx (outil de publication, hors du banc).
"""
import argparse
import datetime
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

BILLING = Path(__file__).resolve().parents[4] / "benchmarks" / "billing"
sys.path.insert(0, str(BILLING))
from bench import hypotheses, reponses, tables, tables_reel  # noqa: E402

KEYS = ("none", "attempt", "business")
FAULTS = ("none", "F1", "F2", "F3", "F4", "F4e", "F5", "F6", "F7", "F8")
GREY = RGBColor(0x40, 0x40, 0x40)


def styled(doc):
    for name, size in (("Normal", 10.5), ("Heading 1", 15), ("Heading 2", 12.5), ("Title", 22)):
        st = doc.styles[name]
        st.font.name = "Liberation Sans"
        st.font.size = Pt(size)
        if name != "Normal":
            st.font.color.rgb = RGBColor(0, 0, 0)
    return doc


def table(doc, header, rows, size=8):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = str(text)
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = "" if text is None else str(text)
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(size)
    doc.add_paragraph()
    return t


def note(doc, text):
    par = doc.add_paragraph(text)
    for run in par.runs:
        run.font.color.rgb = GREY
        run.font.size = Pt(9)


def provenance(doc, campaign, label):
    start = next((c for c in campaign if c.get("event") == "start"), {})
    end = next((c for c in campaign if c.get("event") in ("end", "stopped", "interrupted")), None)
    table(doc, ["Élément", label], [
        ["Commit du dépôt", f"{(start.get('git_commit') or '')[:12]}"
                            + (" (fichiers non commités présents)" if start.get("git_dirty") else "")],
        ["Empreinte du banc", (start.get("bench_sha256") or "")[:16]],
        ["Empreinte du binaire swarm", (start.get("swarm_sha256") or "")[:16]],
        ["Fin", end.get("event") if end else "campagne en cours : résultats provisoires"],
        ["Avertissement", (end or {}).get("warning") or (end or {}).get("reason") or "aucun"],
    ])
    return end is not None


def scripted(doc, path):
    records, campaign = tables.load(path)
    doc.add_heading("1. Campagne scriptée", level=1)
    complete = provenance(doc, campaign, Path(path).name)
    status = Counter((r["condition"], r["status"]) for r in records)
    table(doc, ["Condition", "OK", "INVALIDE", "DÉLAI", "ERREUR"],
          [[c] + [status.get((c, s), 0) for s in ("OK", "INVALIDE", "DÉLAI", "ERREUR")] for c in ("B0", "B1", "S")])
    summary = tables.summarize(records)

    doc.add_heading("1.1 Synthèse par case", level=2)
    note(doc, "Chaque case : exécutions présentant au moins un doublon / un paiement inexact / un impayé / un "
              "faux succès, sur les exécutions valides (n). Les mesures ne s'additionnent pas.")
    rows = []
    for c in ("B0", "B1", "S"):
        for k in KEYS:
            row = [c, k]
            for f in FAULTS:
                s = summary.get((c, k, f))
                row.append("—" if not s else f"{s['double_runs']}/{s['wrong_runs']}/{s['unpaid_runs']}/"
                                             f"{s['false_success_runs']} (n={s['ok']})")
            rows.append(row)
    table(doc, ["Cond.", "Clé", *FAULTS], rows, size=6.5)

    flagged = [(c, s) for c, s in summary.items() if s["flagged"]]
    doc.add_heading("1.2 Exclusions", level=2)
    excluded = [[*c, s["invalid"], s["timeout"], s["error"], f"{s['excluded']}/{s['total']}",
                 "signalée" if s["flagged"] else ""] for c, s in summary.items() if s["excluded"]]
    if excluded:
        table(doc, ["Cond.", "Clé", "Faute", "INVALIDE", "DÉLAI", "ERREUR", "Exclues", "> 10 %"], excluded)
    else:
        doc.add_paragraph("Aucune exécution exclue.")
    if flagged:
        note(doc, "Cases au-delà de 10 % d'exclusions : signalées et discutées, jamais retirées (protocole 6.7).")

    doc.add_heading("1.3 Hypothèses H1 à H6", level=2)
    note(doc, "Verdicts calculés par bench/hypotheses.py selon des règles fixées avant la fin de la campagne.")
    for name, (verdict, details) in hypotheses.verdicts(records).items():
        doc.add_paragraph(f"{name} : {verdict}", style="List Bullet")
        gaps = [d for d in details if "ÉCART" in d or "aucune" in d or "rapport" in d or "S/F6" in d or "étiquetage" in d]
        for d in gaps:
            doc.add_paragraph(d, style="List Bullet 2")

    doc.add_heading("1.4 Auteur de l'arrêt sous F4 et F4e (S)", level=2)
    table(doc, ["Cond.", "Clé", "Faute", "Moteur", "Règlement", "Aucun"],
          [[*c, v["engine"], v["settlement"], v["none"]] for c, v in tables.stopped_by(records).items()])
    doc.add_heading("1.5 F5 absorbée", level=2)
    table(doc, ["Cond.", "Clé", "Valides", "Absorbées"],
          [[c[0], c[1], v["ok"], v["absorbed"]] for c, v in tables.f5_absorbed(records).items()])
    return complete


def answers(doc, runs, path):
    doc.add_heading("2.5 Réponses finales des agents B0-réel (lecture a posteriori)", level=2)
    note(doc, "Lecture ajoutée après observation (amendement du 6 octobre 2026 au soir). Elle complète la mesure "
              "pré-enregistrée du succès déclaré (sortie normale de l'agent) sans la remplacer. Classes par règles "
              "lexicales fixes (bench/reponses.py) ; textes intégraux en annexe.")
    final = reponses.final_answers(path)
    pre = {(r["fault"], r["seed"]): r for r in runs if r["condition"] == "B0r"}
    table(doc, ["Faute", "Graine", "Payées", "Impayées", "Faux succès (pré-enregistré)", "Réponse finale (a posteriori)"],
          [[f, s, (pre[(f, s)].get("metrics") or {}).get("payments"), (pre[(f, s)].get("metrics") or {}).get("unpaid"),
            "oui" if pre[(f, s)].get("false_success") else "non", a["class"]]
           for (f, s), a in sorted(final.items()) if (f, s) in pre])
    return final


def appendix(doc, final):
    doc.add_heading("Annexe : réponses finales intégrales des agents B0-réel", level=1)
    for (f, s), a in sorted(final.items()):
        doc.add_paragraph(f"{f}, graine {s}, {a['call']} : {a['class']}", style="List Bullet")
        note(doc, a["result"] or "(vide)")


def real(doc, path, responses=None):
    runs, campaign = tables_reel.load(path)
    doc.add_heading("2. Lot à agents réels (exploratoire)", level=1)
    provenance(doc, campaign, Path(path).name)
    note(doc, "B0-réel contre S-réel compare deux architectures complètes ; seul W-réel contre S-réel isole "
              "l'apport du moteur. Responsable et revue de S-réel restent scriptés.")
    rows = [tables_reel.run_row(r) for r in runs]
    scen = tables_reel.scenarios(rows)
    header = ["Faute", "Cond.", "Valides/k", "Correctes", "pass^k", "Faux succès", "IBAN F9 payé",
              "Coût agent ($)", "Sans coût", "Tours méd.", "Durée méd. (s)"]
    for title, faults in (("2.1 Fautes communes", tables_reel.COMMON),
                          ("2.2 Fautes propres à la séparation", tables_reel.SEPARATION)):
        doc.add_heading(title, level=2)
        table(doc, header, [[f, c, f"{s['valid']}/{s['k']}", s["correct"], "oui" if s["pass_k"] else "non",
                             s["false_success"], s["attacker_paid_runs"], f"{s['cost_usd']:.3f}",
                             s["unpriced_calls"], s["median_turns"], s["median_duration_s"]]
                            for (c, f), s in tables_reel.ordered(scen) if f in faults])
    doc.add_heading("2.3 Doublons et clés envoyées (Q8)", level=2)
    doubles = [x for x in rows if x["double_keys"]]
    for x in doubles:
        doc.add_paragraph(f"{x['condition']} {x['fault']} graine {x['seed']} : {x['double_keys']}", style="List Bullet")
    if not doubles:
        doc.add_paragraph("Aucune facture payée plusieurs fois.")
    doc.add_heading("2.4 Exécutions", level=2)
    table(doc, ["Cond.", "Faute", "Graine", "Statut", "Paiements", "Doublons", "Inexacts", "Impayés",
                "Faux succès", "IBAN F9", "Coût ($)", "Tours", "Durée (s)"],
          [[x["condition"], x["fault"], x["seed"], x["status"], x["payments"], x["doubles"], x["wrong"], x["unpaid"],
            "oui" if x["false_success"] else "non", x["attacker_paid"], f"{x['cost_usd']:.3f}", x["turns"],
            x["duration_s"]] for x in sorted(rows, key=lambda x: (x["fault"], x["condition"], x["seed"]))], size=7)
    return answers(doc, runs, responses) if responses else None


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--scripted", required=True)
    p.add_argument("--real")
    p.add_argument("--responses", help="réponses des agents extraites par bench/reponses.py")
    p.add_argument("--out", required=True)
    a = p.parse_args(argv)
    doc = styled(Document())
    doc.add_heading("Banc de facturation : résultats", level=0)
    par = doc.add_paragraph(f"Fabrice Pizzi — généré le {datetime.date.today().isoformat()} à partir des fichiers "
                            "de résultats ; données entièrement synthétiques.")
    par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    complete = scripted(doc, a.scripted)
    final = real(doc, a.real, a.responses) if a.real else None
    doc.add_heading("Limites", level=1)
    for text in ("Agents, responsable et revue scriptés dans la campagne principale : les résultats valent pour le "
                 "mécanisme, pas pour le comportement d'un modèle.",
                 "Un seul scénario métier, une API de paiement simulée qui déduplique parfaitement par clé.",
                 "Lot réel exploratoire : cinq essais par scénario, un modèle, un client.",
                 "H6 et l'attribution des blocages relèvent de l'étiquetage humain (section 7)."):
        doc.add_paragraph(text, style="List Bullet")
    if final:
        appendix(doc, final)
    if not complete:
        note(doc, "Campagne scriptée en cours au moment de la génération : chiffres provisoires.")
    doc.save(a.out)
    print(a.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
