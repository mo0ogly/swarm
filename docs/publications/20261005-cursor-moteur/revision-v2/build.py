from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
B=Path(__file__).parent
D=Document();s=D.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=s.bottom_margin=Inches(.72);s.left_margin=s.right_margin=Inches(.8)
for name,size in [('Normal',11.5),('Title',27),('Subtitle',11),('Heading 1',16),('Heading 2',13)]:
 st=D.styles[name];st.font.name='Liberation Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(8)
 if name=='Normal':st.paragraph_format.line_spacing=1.12
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(14);st.paragraph_format.keep_with_next=True
for el in [D.styles.element,D.element]:
 for e in list(el.iter(qn('w:pBdr'))):e.getparent().remove(e)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=f.add_run('Fabrice Pizzi  ·  ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
content=[]
def p(t):
 paragraph=D.add_paragraph(); content.append(t)
 import re
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 pos=0
 for match in re.finditer(r'https://[^\s]+',t):
  paragraph.add_run(t[pos:match.start()])
  link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(match.group(),RT.HYPERLINK,is_external=True))
  run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'205B83');props.append(color)
  under=OxmlElement('w:u');under.set(qn('w:val'),'single');props.append(under);run.append(props)
  text=OxmlElement('w:t');text.text=match.group();run.append(text);link.append(run);paragraph._p.append(link);pos=match.end()
 paragraph.add_run(t[pos:])
def h(t):D.add_heading(t,level=1);content.append(t)
D.add_paragraph('Orchestrer les agents IA avec un moteur déterministe',style='Title')
D.add_paragraph('Fabrice Pizzi · Article LinkedIn · Version révisée · 5 octobre 2026',style='Subtitle')
p('Je suis parti d’un besoin concret : voir ce que font mes agents IA, organiser leur travail et comprendre pourquoi une mission avance, attend ou recommence. Une réponse convaincante ne me suffit pas. Je veux savoir quelles actions ont été exécutées, sur quelle version et avec quelles preuves.')
p('Les recherches de Cursor sur les bases de code autonomes ont nourri ces travaux. Elles m’ont conduit à une question d’architecture : quelles décisions confier aux agents, et lesquelles faire appliquer par un moteur explicite pendant toute l’exécution ?')
h('Ce que je retiens des recherches de Cursor')
p('Dans Towards self-driving codebases, Wilson Lin décrit une organisation avec un planificateur principal, des sous-planificateurs et des exécutants. Les résultats remontent au responsable du périmètre par une passation structurée. L’article montre aussi les limites de la coordination improvisée et l’importance des traces horodatées.')
p('Un autre enseignement mérite d’être conservé : Cursor a retiré un intégrateur central devenu un goulot. Exiger un code parfaitement correct avant chaque commit ralentissait fortement le travail. Ce retour d’expérience invite à distinguer la progression exploratoire et les conditions d’une livraison.')
p('Le système final de Cursor privilégie la hiérarchie de planification, des exécutants au périmètre limité et les passations. Les contrôles de lancement et la validation explicite que je retiens pour Swarm sont mes choix d’architecture. Je ne les présente pas comme des prescriptions de Cursor.')
h('Pourquoi le contrôle final ne suffit pas')
p('Une revue finale peut découvrir un défaut. Elle arrive trop tard pour empêcher deux agents de modifier le même espace partagé, une reprise en double ou une succession d’appels sans progrès utile. Les ressources ont déjà été consommées, et certains effets ont déjà eu lieu.')
p('C’est pourquoi je place le moteur déterministe au cœur du traitement. Il intervient avant les départs, conserve l’état pendant le travail et contrôle les conditions de reprise. La validation finale reste nécessaire ; elle complète ces contrôles pendant l’exécution.')
D.add_page_break()
h('Le partage des responsabilités')
p('Les agents proposent un plan, explorent, produisent et expliquent leurs résultats. Le moteur applique les droits, les dépendances, les réservations d’espace et les limites configurées. L’humain fixe les objectifs, autorise le périmètre et arbitre les exceptions qui lui reviennent.')
im=Image.open(B/'orchestration.png');w,hpx=im.size;maxw=6.5;maxh=6.15;scale=min(maxw/w,maxh/hpx);D.add_picture(str(B/'orchestration.png'),width=Inches(w*scale))
c=D.add_paragraph('Schéma de principe de mon approche · Gris : évolutions prévues · Violet : agents IA · Bleu : moteur · Vert : preuves · Ambre : décisions');c.paragraph_format.space_after=Pt(10)
for r in c.runs:r.font.size=Pt(9)
p('Une proposition du modèle ne devient pas automatiquement une action autorisée. Un processus terminé ne devient pas automatiquement un résultat accepté. Ces deux distinctions rendent le fonctionnement plus compréhensible et les décisions vérifiables.')
D.add_page_break()
h('Ce que cette combinaison apporte')
p('D’abord, une coordination explicite. Les dépendances indiquent ce qui doit être validé avant de poursuivre. Les responsabilités identifient qui produit, qui examine et qui décide. Le journal relie les actions à leurs tentatives, au lieu de laisser l’utilisateur interpréter un simple statut « en cours ».')
p('Ensuite, une meilleure lecture de la consommation. Je distingue les appels d’outils, les requêtes au modèle, les jetons, la durée et le coût rapporté. Ces unités ne sont pas interchangeables. Un compteur d’outils ne permet pas de calculer une facture, et un coût absent doit rester inconnu.')
p('Cette instrumentation aide à poser les bonnes questions : l’agent explore-t-il utilement ? Réexécute-t-il les mêmes contrôles ? Le moteur a-t-il provoqué une reprise inutile ? L’environnement manque-t-il d’une capacité ? Je peux alors chercher à optimiser le chemin suivi, au-delà du résultat annoncé.')
p('Enfin, la reprise devient une décision justifiée. Si un contrôle échoue parce que le navigateur est indisponible, relancer la production du code ne corrige pas la cause. Il faut vérifier la capacité de l’environnement puis reprendre la validation pertinente, en conservant la trace de l’échec précédent.')
h('Rendre les décisions compréhensibles dans l’interface')
p('Le plan que je viens de formaliser étend cette approche à l’édition du graphe et à l’automatisation. Ces évolutions sont prévues, pas encore livrées : une modification du graphe prépare un brouillon, le moteur en expose les conséquences, puis vérifie à nouveau les conditions au moment de l’application.')
p('Déplacer une carte ne doit pas invalider une preuve. Modifier un prérequis ou le candidat testé peut, au contraire, la rendre périmée. Une programmation doit soumettre une demande durable aux mêmes règles que le lancement manuel. Après une réponse perdue, répéter cette demande ne doit pas démarrer deux agents.')
p('Je veux que ces opérations soient accessibles depuis le web et la CLI, avec la même décision métier. Une attente doit expliquer sa cause, qui peut agir et ce qui permettra la reprise. Les rôles, le travail en cours et les résultats validés doivent rester visuellement distincts.')
D.add_page_break()
h('Le moteur peut lui aussi créer le problème')
p('Déterministe ne signifie pas correct. Une règle erronée sera appliquée avec régularité. Un plafond mal calibré peut interrompre un travail utile ; une règle de fraîcheur trop large peut invalider une preuve après un changement purement administratif. Ajouter des contrôles sans tester leurs effets peut rendre le système moins efficace.')
p('Mon objectif est donc un contrôle proportionné : autoriser les actions dans un périmètre clair, vérifier ce qui a réellement changé et éviter les reprises identiques. Je conserve les contrôles ciblés près du travail, puis une validation d’ensemble avant livraison. Je ne revendique pas encore un gain général de coût ou de performance : il faut le mesurer sur des missions comparables.')
h('La question que je veux partager')
p('Construire des agents autonomes, c’est aussi concevoir le système qui leur permet d’agir sans perdre la responsabilité du résultat. Je veux garder leur capacité d’exploration tout en rendant leurs actions, leurs limites et leurs preuves observables.')
p('Dans vos projets, où placez-vous ces contrôles : avant les actions, pendant l’exécution, à la livraison ? Et comment distinguez-vous une erreur de l’agent d’un défaut du moteur qui le pilote ?')
p('Swarm : https://github.com/mo0ogly/swarm')
p('Recherche à l’origine de cette réflexion : Wilson Lin, Towards self-driving codebases, Cursor, 5 février 2026 — https://cursor.com/blog/self-driving-codebases')
p('#AgenticAI #SoftwareEngineering #MultiAgentSystems #Observability')
D.core_properties.author='Fabrice Pizzi';D.core_properties.title='Orchestrer les agents IA avec un moteur déterministe';D.core_properties.subject='Article LinkedIn inspiré des recherches de Cursor';D.save(B/'Article_LinkedIn_Cursor_moteur_deterministe_v2.docx')
(B/'article.txt').write_text('\n\n'.join(content)+'\n')
print('DOCX créé',len(' '.join(content).split()),'mots')
