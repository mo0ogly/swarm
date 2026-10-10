#!/usr/bin/env python3
"""Dossier de recherche en Word : assemble les chapitres de ce répertoire par pandoc.

Usage : build.py SORTIE.docx
Les chapitres sont pris dans l'ordre de leur nom (01-… à 95-…).
"""
import subprocess
import sys
import tempfile
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
HEAD = ["% Médiation déterministe des agents LLM sur processus à effets irréversibles",
        "% Fabrice Pizzi", "% Dossier de recherche, base de thèse — 7 octobre 2026, révision du 8 octobre 2026", ""]


def main(out):
    chapters = sorted(p for p in HERE.glob("[0-9][0-9]-*.md"))
    contents = ["# Plan du dossier", ""]
    for p in chapters:
        contents.extend(line[2:] for line in p.read_text().splitlines() if line.startswith("# "))
        contents.append("")
    text = "\n".join(HEAD) + "\n\n" + "\n\n".join(contents) + "\n\n" + "\n\n".join(p.read_text().strip() for p in chapters)
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False) as f:
        f.write(text)
        source = f.name
    try:
        subprocess.run(["pandoc", source, "-o", out, "--metadata=lang:fr-FR",
                        "--from", "markdown+pipe_tables-yaml_metadata_block-multiline_tables-simple_tables"], check=True)
    finally:
        Path(source).unlink(missing_ok=True)
    # Stable visible contents instead of an empty, unrefreshed Word TOC field.
    doc = Document(out)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2.5)
        section.top_margin = section.bottom_margin = Cm(2.5)
    title = doc.styles["Title"]
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.font.size = Pt(20)
    title.font.bold = True
    for paragraph in doc.paragraphs:
        if paragraph.style.name in ("Title", "Subtitle"):
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
    # Split the wide four-metric results table into two readable panels.
    for table in list(doc.tables):
        if len(table.columns) > 8:
            for indices in ([0, 1, 2, 3, 4, 5, 6], [0, 1, 7, 8, 9, 10]):
                panel = deepcopy(table._tbl)
                for row in panel.findall(qn("w:tr")):
                    for index, cell in reversed(list(enumerate(row.findall(qn("w:tc"))))):
                        if index not in indices:
                            row.remove(cell)
                grid = panel.find(qn("w:tblGrid"))
                for index, col in reversed(list(enumerate(list(grid)))):
                    if index not in indices:
                        grid.remove(col)
                table._tbl.addprevious(panel)
            table._tbl.getparent().remove(table._tbl)
    available = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    for table in doc.tables:
        table.autofit = False
        weights = [1] * len(table.columns)
        if [cell.text for cell in table.rows[0].cells] == ["Invariant", "Test", "Fichier", "État git"]:
            weights = [1.25, 4.25, 3.3, 1.2]
        for col, weight in zip(table.columns, weights):
            col.width = int(available * weight / sum(weights))
        for row in table.rows:
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            for cell, weight in zip(row.cells, weights):
                cell.width = int(available * weight / sum(weights))
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        parts = run.text.split("/")
                        if len(parts) == 4 and all(part.isdigit() for part in parts):
                            run.text = "/".join(parts[:2]) + "/\n" + "/".join(parts[2:])
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1])
