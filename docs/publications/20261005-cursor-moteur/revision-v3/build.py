from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
B=Path(__file__).parent
D=Document();s=D.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=s.bottom_margin=Inches(.72);s.left_margin=s.right_margin=Inches(.8)
for name,size in [('Normal',11.5),('Title',27),('Subtitle',11),('Heading 1',16),('Heading 2',13)]:
 st=D.styles[name];st.font.name='Liberation Sans';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(8)
 if name=='Normal':st.paragraph_format.line_spacing=1.12
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(14);st.paragraph_format.keep_with_next=True
for el in [D.styles.element,D.element]:
 for e in list(el.iter(qn('w:pBdr'))):e.getparent().remove(e)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=f.add_run('Fabrice Pizzi  ·  ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
content=[]
def p(t):
 paragraph=D.add_paragraph(); content.append(t)
 import re
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 pos=0
 for match in re.finditer(r'https://[^\s]+',t):
  paragraph.add_run(t[pos:match.start()])
  link=OxmlElement('w:hyperlink');link.set(qn('r:id'),paragraph.part.relate_to(match.group(),RT.HYPERLINK,is_external=True))
  run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'205B83');props.append(color)
  under=OxmlElement('w:u');under.set(qn('w:val'),'single');props.append(under);run.append(props)
  text=OxmlElement('w:t');text.text=match.group();run.append(text);link.append(run);paragraph._p.append(link);pos=match.end()
 paragraph.add_run(t[pos:])
def h(t):D.add_heading(t,level=1);content.append(t)
def fig(name,caption,maxh=6.15):
 im=Image.open(B/name);w,hpx=im.size;scale=min(6.5/w,maxh/hpx);D.add_picture(str(B/name),width=Inches(w*scale))
 c=D.add_paragraph(caption);c.paragraph_format.space_after=Pt(10)
 for r in c.runs:r.font.size=Pt(9)
D.add_paragraph('Agents IA et processus critiques : pourquoi un moteur déterministe',style='Title')
D.add_paragraph('Fabrice Pizzi · Article LinkedIn · Version 3 · 5 octobre 2026',style='Subtitle')
p('Je travaille sur une question précise : peut-on confier à des agents IA une partie d’un processus métier où une erreur coûte de l’argent, comme la facturation ou une transaction bancaire ? Sur ce terrain, une réponse convaincante ne suffit pas. Il faut savoir quelles actions ont été exécutées, sur quelle version, avec quelles preuves, et surtout empêcher qu’un effet irréversible parte deux fois ou parte faux.')
p('Les recherches de Cursor sur les bases de code autonomes ont nourri ces travaux. Elles m’ont conduit à une question d’architecture : quelles décisions confier aux agents, et lesquelles faire appliquer par un moteur explicite pendant toute l’exécution ?')
h('Ce que je retiens des recherches de Cursor')
p('Dans Towards self-driving codebases, Wilson Lin décrit une organisation avec un planificateur principal, des sous-planificateurs et des exécutants. Les résultats remontent au responsable du périmètre par une passation structurée. L’article montre aussi les limites de la coordination improvisée et l’importance des traces horodatées.')
p('Un autre enseignement mérite d’être conservé : Cursor a retiré un intégrateur central devenu un goulot. Exiger un code parfaitement correct avant chaque commit ralentissait fortement le travail. Pour du code, qu’un autre agent corrigera bientôt, cette tolérance se défend. Pour un virement, elle ne se défend plus : l’erreur ne se corrige pas au commit suivant.')
p('Les contrôles de lancement et la validation explicite que je retiens pour Swarm sont mes choix d’architecture. Je ne les présente pas comme des prescriptions de Cursor.')
h('Pourquoi le contrôle final ne suffit pas')
p('Une revue finale peut découvrir un défaut. Elle arrive trop tard pour empêcher deux agents de travailler sur le même espace, une reprise en double ou un paiement déjà parti. Les ressources ont été consommées, et l’effet a eu lieu.')
p('C’est pourquoi je place le moteur déterministe au cœur du traitement. Il intervient avant les départs, conserve l’état pendant le travail et contrôle les conditions de reprise. La validation finale reste nécessaire ; elle complète ces contrôles pendant l’exécution.')
h('Un cas concret : le double paiement')
p('Un agent prépare un lot de factures fournisseurs. Le règlement envoie un paiement, la banque l’exécute, mais la réponse se perd sur le réseau. Le règlement réessaie. Sans précaution, deux virements partent.')
p('La première protection est connue : une clé d’idempotence, que l’API de paiement utilise pour reconnaître une demande déjà traitée. Encore faut-il choisir la bonne clé. Une clé qui dépend de la tentative laisse passer le doublon : la tentative suivante produit une clé neuve, et la banque paie à nouveau. Seule une clé métier, fournisseur et numéro de facture, reconnaît la même intention.')
p('Le problème est documenté : les auteurs d’ACRFence (Zheng et al., 2026) ont montré qu’un agent LLM qui reprend après un incident réécrit légèrement sa requête, ce qui met en défaut les protections fondées sur des requêtes identiques ; ils ont obtenu un paiement en double dans chacun de leurs dix essais avec reprise. https://arxiv.org/abs/2603.20625')
fig('double_paiement.png','Réponse perdue puis nouvel essai : seule la clé métier empêche le second virement',maxh=5.6)
p('Cette clé ne règle pourtant qu’une partie du problème. Elle empêche la répétition à l’identique. Elle n’empêche pas un paiement correct en apparence mais faux sur le fond : un montant modifié après validation, un lot remplacé par une version plus ancienne, un lot partiel réglé parce qu’un agent a épuisé son budget. C’est là que le moteur doit intervenir, avant l’effet et non après.')
D.add_page_break()
h('Le partage des responsabilités')
p('Les agents proposent, explorent, produisent et expliquent. Dans mon scénario, l’agent prépare le lot de paiements ; il n’a aucun droit de payer. Le moteur applique les droits, les dépendances et les limites, et lie chaque validation au lot exact qu’elle a examiné. Un règlement déterministe n’exécute que ce lot accepté. L’humain fixe les objectifs, autorise le périmètre et arbitre les exceptions.')
fig('chemin_paiement.png','Chemin d’un paiement sous contrôle du moteur · Violet : agent IA · Bleu : moteur et règlement · Vert : effets et mesures · Ambre : décisions humaines · Rouge : refus sans effet')
p('Une proposition du modèle ne devient pas automatiquement une action autorisée. Un processus terminé ne devient pas automatiquement un résultat accepté. Dans le monde bancaire, ce principe porte un nom ancien : le contrôle à quatre yeux. Le moteur en est la version exécutable.')
D.add_page_break()
h('Ce que cette combinaison apporte')
fig('orchestration.png','Schéma de principe de l’orchestration · Gris : évolutions prévues · Violet : agents IA · Bleu : moteur · Vert : preuves · Ambre : décisions')
p('D’abord, une coordination explicite. Les dépendances indiquent ce qui doit être validé avant de poursuivre. Le journal relie les actions à leurs tentatives, au lieu de laisser l’utilisateur interpréter un simple statut « en cours ».')
p('Ensuite, une lecture honnête de la consommation. Je distingue les appels d’outils, les requêtes au modèle, les jetons, la durée et le coût rapporté. Ces unités ne sont pas interchangeables, et un coût absent doit rester inconnu.')
p('Enfin, la reprise devient une décision justifiée. Si un contrôle échoue parce que le service est indisponible, régénérer le lot ne corrige pas la cause. Il faut rétablir l’environnement puis reprendre la validation pertinente, en conservant la trace de l’échec précédent.')
h('Le moteur peut lui aussi créer le problème')
p('Déterministe ne signifie pas correct. Une règle erronée sera appliquée avec régularité. Un plafond mal calibré peut interrompre un travail utile ; une règle de fraîcheur trop large peut bloquer un lot valide. Un contrôle de trop a lui aussi un coût. L’objectif est un contrôle proportionné, et il doit se mesurer.')
h('Ce que je mesure maintenant')
p('Je construis un banc de facturation simulé, sur données entièrement synthétiques : un grand livre en partie double, une API de paiement fictive, des factures fournisseurs. J’y compare trois situations : une chaîne sans moteur avec une revue finale, un contrôle appel par appel, et le moteur complet.')
p('J’y injecte huit fautes reproductibles : réponse perdue, deux lanceurs concurrents, crash entre le paiement et son enregistrement, lot modifié après validation, service indisponible, budget épuisé, responsable de mission figé puis réveillé, et rapport ancien relayé après une nouvelle tentative. Chaque faute est croisée avec trois choix de clé d’idempotence.')
p('Une règle de méthode : chaque faute doit d’abord produire le défaut attendu dans la chaîne sans protection. Sinon, l’injection ne prouve rien, et un zéro du côté du moteur ne vaudrait rien non plus. Les mesures se lisent dans le grand livre, jamais dans le discours des agents.')
p('Je ne publie aucun résultat aujourd’hui : la campagne n’a pas encore tourné. Je reviendrai avec les chiffres, y compris ceux qui ne vont pas dans mon sens, en particulier les blocages à tort imputables au moteur.')
h('La question que je veux partager')
p('Construire des agents autonomes, c’est aussi concevoir le système qui leur permet d’agir sans perdre la responsabilité du résultat. Plus l’effet est irréversible, plus cette responsabilité doit être portée par un mécanisme vérifiable plutôt que par la qualité d’une réponse.')
p('Dans vos projets, où placez-vous ces contrôles : avant les actions, pendant l’exécution, à la livraison ? Et pour ceux qui travaillent sur des flux financiers : quelle faute vous inquiète le plus, et laquelle votre architecture actuelle ne verrait qu’après coup ?')
p('Swarm : https://github.com/mo0ogly/swarm')
p('Recherche à l’origine de cette réflexion : Wilson Lin, Towards self-driving codebases, Cursor, 5 février 2026 — https://cursor.com/blog/self-driving-codebases')
p('#AgenticAI #FinTech #MultiAgentSystems #SoftwareEngineering #Observability')
D.core_properties.author='Fabrice Pizzi';D.core_properties.title='Agents IA et processus critiques : pourquoi un moteur déterministe';D.core_properties.subject='Article LinkedIn, version 3 : processus financiers et banc de facturation';D.save(B/'Article_LinkedIn_Cursor_moteur_deterministe_v3.docx')
(B/'article.txt').write_text('\n\n'.join(content)+'\n')
print('DOCX créé',len(' '.join(content).split()),'mots')
